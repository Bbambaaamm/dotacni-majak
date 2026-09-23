from __future__ import annotations

import hashlib
import re
from datetime import datetime, time, timezone
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from urllib.parse import urljoin, urlsplit
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from dotacni_majak_source_sdk import (
    AdapterContext,
    ArtifactFetchResult,
    AuthorityLevel,
    DiscoveryItem,
    DiscoveryPage,
    FetchState,
    FetchValidators,
    HealthReport,
    HealthStatus,
    NativeRecord,
    RecordFetchResult,
    RemoteArtifactRef,
    RetrievalMode,
    SourceAdapter,
    SourceCheckpoint,
    SourceDescriptor,
)


LISTING_URL = "https://irop.gov.cz/cs/vyzvy-2021-2027"
_ALLOWED_HOSTS = {"irop.gov.cz", "www.irop.gov.cz"}
_LOCAL_TZ = ZoneInfo("Europe/Prague")


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def _is_detail_url(url: str) -> bool:
    parsed = urlsplit(url)
    path = parsed.path.casefold().rstrip("/")
    return (
        parsed.hostname in _ALLOWED_HOSTS
        and "/vyzvy-2021-2027/vyzvy/" in path
        and path != "/cs/vyzvy-2021-2027"
    )


def _external_id(title: str, url: str) -> str:
    match = re.search(r"\b(\d{1,3})\.\s*výzva\s+irop\b", title, re.I)
    if match:
        return f"IROP-{int(match.group(1))}"
    path_match = re.search(r"/(\d{1,3})vyzvairop/?$", urlsplit(url).path, re.I)
    if path_match:
        return f"IROP-{int(path_match.group(1))}"
    return "IROP-" + hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]


def _parse_date(value: str | None, *, end_of_day: bool = False) -> datetime | None:
    if not value:
        return None
    cleaned = _clean(value)
    for fmt in ("%d. %m. %Y", "%d.%m.%Y", "%d. %m. %Y %H:%M", "%d.%m.%Y %H:%M"):
        try:
            parsed = datetime.strptime(cleaned, fmt)
            if end_of_day and "%H" not in fmt:
                parsed = datetime.combine(parsed.date(), time.max)
            return parsed.replace(tzinfo=_LOCAL_TZ).astimezone(timezone.utc)
        except ValueError:
            continue
    return None


def _status(value: str | None) -> str:
    lowered = (value or "").casefold()
    if "zruš" in lowered:
        return "CANCELLED"
    if "pozastav" in lowered:
        return "PAUSED"
    if "otev" in lowered:
        return "OPEN"
    if "plán" in lowered:
        return "PLANNED"
    if "uzav" in lowered or "ukon" in lowered:
        return "CLOSED"
    return "UNKNOWN"


def _table_pairs(soup: BeautifulSoup) -> dict[str, str]:
    pairs: dict[str, str] = {}
    for row in soup.find_all("tr"):
        cells = [
            _clean(cell.get_text(" ", strip=True))
            for cell in row.find_all(["th", "td"])
        ]
        if len(cells) >= 2 and cells[0]:
            pairs[cells[0].rstrip(":").casefold()] = cells[1]
    return pairs


def _pair(pairs: dict[str, str], label: str) -> str | None:
    return pairs.get(label.rstrip(":").casefold())


def _money_czk(value: str | None) -> int | None:
    if not value:
        return None
    normalized = value.replace("\xa0", " ")
    direct = re.search(r"([0-9][0-9\s]*)\s*Kč\b", normalized, re.I)
    if direct:
        return int(re.sub(r"\s+", "", direct.group(1)))
    billion = re.search(r"([0-9]+(?:[,.][0-9]+)?)\s*mld\.?\s*Kč", normalized, re.I)
    if billion:
        return int(round(float(billion.group(1).replace(",", ".")) * 1_000_000_000))
    million = re.search(r"([0-9]+(?:[,.][0-9]+)?)\s*mil\.?\s*Kč", normalized, re.I)
    if million:
        return int(round(float(million.group(1).replace(",", ".")) * 1_000_000))
    return None


def _section_text(soup: BeautifulSoup, heading_text: str) -> str | None:
    target = heading_text.casefold()
    for heading in soup.find_all(re.compile(r"^h[1-6]$")):
        if _clean(heading.get_text(" ", strip=True)).casefold() != target:
            continue
        parts: list[str] = []
        for element in heading.find_all_next():
            if element is not heading and element.name and re.fullmatch(r"h[1-6]", element.name):
                break
            if element.name not in {"p", "li", "tr"}:
                continue
            value = _clean(element.get_text(" ", strip=True))
            if value and value not in parts:
                parts.append(value)
        return "\n".join(parts) or None
    return None


def _history(soup: BeautifulSoup) -> list[dict[str, str]]:
    heading = None
    for candidate in soup.find_all(re.compile(r"^h[1-6]$")):
        if _clean(candidate.get_text(" ", strip=True)).casefold() == "aktuality":
            heading = candidate
            break
    if heading is None:
        return []

    events: list[dict[str, str]] = []
    for element in heading.find_all_next():
        if element is not heading and element.name and re.fullmatch(r"h[1-6]", element.name):
            break
        if element.name not in {"p", "tr"}:
            continue
        text = _clean(element.get_text(" ", strip=True))
        if not text:
            continue
        match = re.match(r"(\d{1,2}\.\s*\d{1,2}\.\s*20\d{2})\s*[|–-]?\s*(.+)", text)
        if match:
            event = {"date": _clean(match.group(1)), "change": _clean(match.group(2))}
            if event not in events:
                events.append(event)
    return events[:100]


def _listing_html_items(html: str) -> list[DiscoveryItem]:
    soup = BeautifulSoup(html, "html.parser")
    items: dict[str, DiscoveryItem] = {}
    for anchor in soup.find_all("a", href=True):
        url = urljoin(LISTING_URL, anchor["href"])
        if not _is_detail_url(url):
            continue
        text = _clean(anchor.get_text(" ", strip=True))
        if not text:
            continue

        title_match = re.match(
            r"(.+?)\s+Příjem\s+žádostí:\s*"
            r"(\d{1,2}\.\s*\d{1,2}\.\s*20\d{2})\s*[-–]\s*"
            r"(\d{1,2}\.\s*\d{1,2}\.\s*20\d{2})\s+"
            r"(Otevřená|Uzavřená|Plánovaná)",
            text,
            re.I,
        )
        title = _clean(title_match.group(1)) if title_match else text
        metadata: dict[str, Any] = {"discovery_method": "HTML"}
        status_hint = None
        if title_match:
            opens = _parse_date(title_match.group(2))
            closes = _parse_date(title_match.group(3), end_of_day=True)
            status_hint = _status(title_match.group(4))
            metadata.update(
                {
                    "submissionOpenAt": opens.isoformat() if opens else None,
                    "submissionCloseAt": closes.isoformat() if closes else None,
                }
            )

        external_id = _external_id(title, url)
        items[external_id] = DiscoveryItem(
            external_id=external_id,
            detail_url=url,
            title_hint=title,
            native_status_hint=status_hint,
            metadata=metadata,
        )
    return sorted(items.values(), key=lambda item: item.external_id)


def _calendar_link(html: str) -> str | None:
    soup = BeautifulSoup(html, "html.parser")
    for anchor in soup.find_all("a", href=True):
        text = _clean(anchor.get_text(" ", strip=True)).casefold()
        href = urljoin(LISTING_URL, anchor["href"])
        if "stáhnout kalendář" in text or "stahnout kalendar" in text:
            return href
        if href.casefold().endswith(".xlsx") and "kalend" in text:
            return href
    return None


def _xlsx_items(content: bytes) -> list[DiscoveryItem]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    items: dict[str, DiscoveryItem] = {}
    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                title_hint: str | None = None
                for cell in row:
                    if cell.value is not None and title_hint is None:
                        title_hint = _clean(str(cell.value))
                    target = cell.hyperlink.target if cell.hyperlink is not None else None
                    if not target and isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                        target = cell.value
                    if not target:
                        continue
                    url = urljoin(LISTING_URL, target)
                    if not _is_detail_url(url):
                        continue
                    title = title_hint or PurePosixPath(urlsplit(url).path).name
                    external_id = _external_id(title, url)
                    items[external_id] = DiscoveryItem(
                        external_id=external_id,
                        detail_url=url,
                        title_hint=title,
                        metadata={"discovery_method": "XLSX"},
                    )
    finally:
        workbook.close()
    return sorted(items.values(), key=lambda item: item.external_id)


def _mime_hint(url: str) -> str | None:
    path = urlsplit(url).path.casefold()
    if path.endswith(".pdf"):
        return "application/pdf"
    if path.endswith(".zip"):
        return "application/zip"
    if path.endswith(".docx"):
        return "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    if path.endswith(".xlsx"):
        return "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return None


def _artifact_role(title: str, call_number: str | None) -> str:
    lowered = title.casefold()
    if "text" in lowered and "výzv" in lowered:
        return "CALL_DOCUMENT"
    if call_number and call_number in lowered and "výzv" in lowered and "pravidla" not in lowered:
        return "CALL_DOCUMENT"
    if "pravidla" in lowered or "postup pro podání žádosti" in lowered:
        return "GUIDELINES"
    if "žádost" in lowered and ("formulář" in lowered or "vzor" in lowered):
        return "APPLICATION_FORM"
    return "ANNEX"


def _artifacts(soup: BeautifulSoup, base_url: str, call_number: str | None) -> list[RemoteArtifactRef]:
    heading = None
    for candidate in soup.find_all(re.compile(r"^h[1-6]$")):
        if _clean(candidate.get_text(" ", strip=True)).casefold() == "připojené soubory":
            heading = candidate
            break
    if heading is None:
        return []

    artifacts: list[RemoteArtifactRef] = []
    seen: set[str] = set()
    for element in heading.find_all_next():
        if element is not heading and element.name and re.fullmatch(r"h[1-6]", element.name):
            break
        if element.name != "a" or not element.get("href"):
            continue
        url = urljoin(base_url, element["href"])
        parsed = urlsplit(url)
        if parsed.hostname not in _ALLOWED_HOSTS:
            continue
        mime = _mime_hint(url)
        if mime is None or url in seen:
            continue
        seen.add(url)

        title = _clean(element.get_text(" ", strip=True)) or PurePosixPath(parsed.path).name
        role = _artifact_role(title, call_number)
        artifacts.append(
            RemoteArtifactRef(
                external_id=hashlib.sha256(url.encode("utf-8")).hexdigest()[:24],
                url=url,
                role=role,
                title=title,
                mime_hint=mime,
                required=(role == "CALL_DOCUMENT"),
            )
        )
    return artifacts


class IropAdapter(SourceAdapter):
    descriptor = SourceDescriptor(
        code="IROP",
        name="Integrovaný regionální operační program 2021–2027",
        authority=AuthorityLevel.OFFICIAL,
        country_code="CZ",
        base_url="https://irop.gov.cz/",
        retrieval_modes=[RetrievalMode.HTML, RetrievalMode.XLSX, RetrievalMode.PDF, RetrievalMode.DOCX],
        allowed_hosts=sorted(_ALLOWED_HOSTS),
        normal_refresh_minutes=240,
        max_concurrency=2,
        requests_per_second=0.5,
        disappearance_confirmation_runs=3,
        adapter_version="0.1.0",
    )

    async def healthcheck(self, ctx: AdapterContext) -> HealthReport:
        started = datetime.now(timezone.utc)
        try:
            response = await ctx.http.get(LISTING_URL)
            body = response.text.casefold()
            healthy = response.status_code == 200 and "výzvy irop 2021-2027" in body and "příjem žádostí" in body
            status = HealthStatus.HEALTHY if healthy else HealthStatus.DEGRADED
            detail = None if healthy else f"Unexpected IROP listing: HTTP {response.status_code}"
        except Exception as exc:
            status = HealthStatus.UNAVAILABLE
            detail = f"{type(exc).__name__}: {exc}"
        elapsed = datetime.now(timezone.utc) - started
        return HealthReport(status=status, checked_at=ctx.now, latency_ms=max(0, int(elapsed.total_seconds() * 1000)), detail=detail)

    async def discover(self, ctx: AdapterContext, checkpoint: SourceCheckpoint | None) -> DiscoveryPage:
        del checkpoint
        response = await ctx.http.get(LISTING_URL)
        if response.status_code >= 400:
            raise RuntimeError(f"IROP listing returned HTTP {response.status_code}")

        listing_snapshot = self._snapshot(
            ctx,
            source_url=LISTING_URL,
            content=response.content,
            mime_type=response.headers.get("content-type", "text/html").split(";", 1)[0],
            headers=response.headers,
        )
        items = {
            item.external_id: item.model_copy(
                update={"metadata": {**item.metadata, "listing_snapshot_id": listing_snapshot}}
            )
            for item in _listing_html_items(response.text)
        }

        calendar_url = _calendar_link(response.text)
        if calendar_url is not None:
            calendar_response = await ctx.http.get(calendar_url)
            if calendar_response.status_code < 400:
                calendar_snapshot = self._snapshot(
                    ctx,
                    source_url=calendar_url,
                    content=calendar_response.content,
                    mime_type=calendar_response.headers.get(
                        "content-type",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ).split(";", 1)[0],
                    headers=calendar_response.headers,
                )
                for item in _xlsx_items(calendar_response.content):
                    existing = items.get(item.external_id)
                    metadata = {
                        **(existing.metadata if existing else {}),
                        **item.metadata,
                        "listing_snapshot_id": listing_snapshot,
                        "calendar_snapshot_id": calendar_snapshot,
                        "calendar_url": calendar_url,
                    }
                    items[item.external_id] = item.model_copy(
                        update={
                            "native_status_hint": existing.native_status_hint if existing else None,
                            "metadata": metadata,
                        }
                    )

        values = sorted(items.values(), key=lambda item: item.external_id)
        return DiscoveryPage(items=values, next_checkpoint=None, is_complete=True, total_hint=len(values))

    async def fetch_record(
        self,
        ctx: AdapterContext,
        item: DiscoveryItem,
        validators: FetchValidators | None = None,
    ) -> RecordFetchResult:
        headers: dict[str, str] = {}
        if validators:
            if validators.etag:
                headers["if-none-match"] = validators.etag
            if validators.last_modified:
                headers["if-modified-since"] = validators.last_modified

        response = await ctx.http.get(str(item.detail_url), headers=headers)
        if response.status_code == 304:
            return RecordFetchResult(state=FetchState.NOT_MODIFIED, http_status=304, etag=response.headers.get("etag"), last_modified=response.headers.get("last-modified"))
        if response.status_code == 404:
            return RecordFetchResult(state=FetchState.GONE, http_status=404)
        if response.status_code >= 400:
            raise RuntimeError(f"IROP detail returned HTTP {response.status_code}")

        snapshot_id = self._snapshot(
            ctx,
            source_url=str(item.detail_url),
            content=response.content,
            mime_type=response.headers.get("content-type", "text/html").split(";", 1)[0],
            headers=response.headers,
        )
        soup = BeautifulSoup(response.text, "html.parser")
        h1 = soup.find("h1")
        title = _clean(h1.get_text(" ", strip=True)) if h1 is not None else (item.title_hint or item.external_id)
        pairs = _table_pairs(soup)

        status_text = _pair(pairs, "Stav výzvy")
        call_type = _pair(pairs, "Druh výzvy")
        access_text = _pair(pairs, "Zpřístupnění žádosti o podporu")
        opens_text = _pair(pairs, "Zahájení příjmu žádostí")
        closes_text = _pair(pairs, "Ukončení příjmu žádostí")
        applicants = _pair(pairs, "Oprávnění žadatelé")
        general_info = _section_text(soup, "Obecné informace")

        allocation = None
        for key, value in pairs.items():
            if key.startswith("finanční alokace výzvy"):
                allocation = _money_czk(value)
                break

        call_number_match = re.search(r"\b(\d{1,3})\.\s*výzva", title, re.I)
        call_number = call_number_match.group(1) if call_number_match else None

        record = NativeRecord(
            source_code=self.descriptor.code,
            external_id=item.external_id,
            detail_url=item.detail_url,
            native_title=title,
            native_status=_status(status_text) if status_text else (item.native_status_hint or "UNKNOWN"),
            raw_fields={
                "programme": "IROP 2021–2027",
                "callNumber": call_number,
                "callType": call_type,
                "applicationAccessAt": _parse_date(access_text).isoformat() if _parse_date(access_text) else None,
                "submissionOpenAt": _parse_date(opens_text).isoformat() if _parse_date(opens_text) else item.metadata.get("submissionOpenAt"),
                "submissionCloseAt": _parse_date(closes_text, end_of_day=True).isoformat() if _parse_date(closes_text, end_of_day=True) else item.metadata.get("submissionCloseAt"),
                "eligibleApplicantsText": applicants,
                "generalInfoText": general_info,
                "allocationCzk": allocation,
                "history": _history(soup),
                "discoveryMethod": item.metadata.get("discovery_method"),
            },
            artifacts=_artifacts(soup, str(item.detail_url), call_number),
            snapshot_ids=[snapshot_id],
        )
        return RecordFetchResult(
            state=FetchState.MODIFIED,
            record=record,
            http_status=response.status_code,
            etag=response.headers.get("etag"),
            last_modified=response.headers.get("last-modified"),
        )

    async def fetch_artifact(
        self,
        ctx: AdapterContext,
        artifact: RemoteArtifactRef,
        validators: FetchValidators | None = None,
    ) -> ArtifactFetchResult:
        headers: dict[str, str] = {}
        if validators:
            if validators.etag:
                headers["if-none-match"] = validators.etag
            if validators.last_modified:
                headers["if-modified-since"] = validators.last_modified
        response = await ctx.http.get(str(artifact.url), headers=headers)
        if response.status_code == 304:
            return ArtifactFetchResult(state=FetchState.NOT_MODIFIED, etag=response.headers.get("etag"), last_modified=response.headers.get("last-modified"))
        if response.status_code == 404:
            return ArtifactFetchResult(state=FetchState.GONE)
        if response.status_code >= 400:
            raise RuntimeError(f"IROP artifact returned HTTP {response.status_code}")

        mime = response.headers.get("content-type", artifact.mime_hint or "application/octet-stream").split(";", 1)[0]
        snapshot_id = self._snapshot(
            ctx,
            source_url=str(artifact.url),
            content=response.content,
            mime_type=mime,
            headers=response.headers,
        )
        return ArtifactFetchResult(
            state=FetchState.MODIFIED,
            snapshot_id=snapshot_id,
            sha256=hashlib.sha256(response.content).hexdigest(),
            mime_type=mime,
            size_bytes=len(response.content),
            etag=response.headers.get("etag"),
            last_modified=response.headers.get("last-modified"),
        )

    def _snapshot(
        self,
        ctx: AdapterContext,
        *,
        source_url: str,
        content: bytes,
        mime_type: str,
        headers: Any,
    ) -> str:
        if ctx.snapshots is None:
            raise RuntimeError("IROP adapter requires ctx.snapshots")
        snapshot = ctx.snapshots.put(
            source_code=self.descriptor.code,
            source_url=source_url,
            content=content,
            mime_type=mime_type,
            retrieved_at=ctx.now,
            validators={
                "etag": headers.get("etag") or "",
                "last_modified": headers.get("last-modified") or "",
            },
        )
        return snapshot.snapshot_id
