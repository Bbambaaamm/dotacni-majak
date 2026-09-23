from __future__ import annotations

import hashlib
import re
from datetime import datetime, time, timezone
from io import BytesIO
from pathlib import PurePosixPath
from typing import Any
from urllib.parse import urljoin, urlsplit
from zoneinfo import ZoneInfo

from bs4 import BeautifulSoup
from openpyxl import load_workbook

from dotacni_majak_source_sdk import (
    AdapterContext,
    ArtifactFetchResult,
    AuthorityLevel,
    DiscoveryItem,
    DiscoveryPage,
    FetchState,
    FetchValidators,
    HealthReport,
    HealthStatus,
    NativeRecord,
    RecordFetchResult,
    RemoteArtifactRef,
    RetrievalMode,
    SourceAdapter,
    SourceCheckpoint,
    SourceDescriptor,
)


LISTING_URL = "https://www.dotaceeu.cz/cs/jak-ziskat-dotaci/vyzvy"
_ALLOWED_HOSTS = {"dotaceeu.cz", "www.dotaceeu.cz"}
_DETAIL_PATH_MARKER = "/jak-ziskat-dotaci/vyzvy/obdobi-"
_LOCAL_TZ = ZoneInfo("Europe/Prague")


def _clean(value: str) -> str:
    return " ".join(value.replace("\xa0", " ").split())


def _is_detail_url(url: str) -> bool:
    parsed = urlsplit(url)
    return (
        parsed.hostname in _ALLOWED_HOSTS
        and _DETAIL_PATH_MARKER in parsed.path.casefold()
        and parsed.path.rstrip("/").casefold()
        != "/cs/jak-ziskat-dotaci/vyzvy"
    )


def _external_id(url: str) -> str:
    slug = PurePosixPath(urlsplit(url).path.rstrip("/")).name
    slug = re.sub(r"[^a-z0-9]+", "-", slug.casefold()).strip("-")
    if slug:
        return f"DOTACEEU-{slug[:96]}"
    return "DOTACEEU-" + hashlib.sha256(url.encode("utf-8")).hexdigest()[:20]


def _parse_local_date(value: str | None, *, end_of_day: bool = False) -> datetime | None:
    if not value:
        return None
    cleaned = _clean(value)
    for fmt in ("%d. %m. %Y", "%d.%m.%Y", "%d. %m. %Y %H:%M"):
        try:
            parsed = datetime.strptime(cleaned, fmt)
            if end_of_day and "%H" not in fmt:
                parsed = datetime.combine(parsed.date(), time.max)
            parsed = parsed.replace(tzinfo=_LOCAL_TZ)
            return parsed.astimezone(timezone.utc)
        except ValueError:
            continue
    return None


def _status(value: str | None) -> str:
    lowered = (value or "").casefold()
    if "zruš" in lowered:
        return "CANCELLED"
    if "pozastav" in lowered:
        return "PAUSED"
    if "otev" in lowered:
        return "OPEN"
    if "plán" in lowered or "priprav" in lowered or "připrav" in lowered:
        return "PLANNED"
    if "ukon" in lowered or "uzav" in lowered:
        return "CLOSED"
    return "UNKNOWN"


def _table_pairs(soup: BeautifulSoup) -> dict[str, str]:
    values: dict[str, str] = {}

    for row in soup.find_all("tr"):
        cells = [
            _clean(cell.get_text(" ", strip=True))
            for cell in row.find_all(["th", "td"])
        ]
        if len(cells) >= 2 and cells[0]:
            values[cells[0].rstrip(":").casefold()] = cells[1]

    for dt in soup.find_all("dt"):
        dd = dt.find_next_sibling("dd")
        if dd is None:
            continue
        key = _clean(dt.get_text(" ", strip=True)).rstrip(":").casefold()
        if key:
            values[key] = _clean(dd.get_text(" ", strip=True))

    return values


def _pair(pairs: dict[str, str], label: str) -> str | None:
    return pairs.get(label.rstrip(":").casefold())


def _more_info_url(soup: BeautifulSoup) -> str | None:
    for row in soup.find_all("tr"):
        cells = row.find_all(["th", "td"])
        if len(cells) < 2:
            continue
        label = _clean(cells[0].get_text(" ", strip=True)).casefold()
        if not label.startswith("více informací") and not label.startswith("vice informaci"):
            continue
        anchor = cells[1].find("a", href=True)
        if anchor:
            return urljoin(LISTING_URL, anchor["href"])
    return None


def _history_rows(soup: BeautifulSoup) -> list[dict[str, str]]:
    heading = None
    for candidate in soup.find_all(re.compile(r"^h[1-6]$")):
        if _clean(candidate.get_text(" ", strip=True)).casefold() == "aktuality":
            heading = candidate
            break
    if heading is None:
        return []

    rows: list[dict[str, str]] = []
    for element in heading.find_all_next():
        if element is not heading and element.name and re.fullmatch(r"h[1-6]", element.name):
            break
        if element.name != "tr":
            continue
        cells = [
            _clean(cell.get_text(" ", strip=True))
            for cell in element.find_all(["th", "td"])
        ]
        cells = [cell for cell in cells if cell]
        if len(cells) >= 2:
            rows.append({"date": cells[0], "change": cells[-1]})
    return rows[:100]


def _listing_html_items(html: str) -> list[DiscoveryItem]:
    soup = BeautifulSoup(html, "html.parser")
    by_url: dict[str, DiscoveryItem] = {}

    for anchor in soup.find_all("a", href=True):
        url = urljoin(LISTING_URL, anchor["href"])
        if not _is_detail_url(url):
            continue
        title = _clean(anchor.get_text(" ", strip=True))
        if not title:
            continue
        by_url[url] = DiscoveryItem(
            external_id=_external_id(url),
            detail_url=url,
            title_hint=title,
            metadata={"discovery_method": "HTML"},
        )

    return sorted(by_url.values(), key=lambda item: item.external_id)


def _xlsx_link(html: str) -> str | None:
    soup = BeautifulSoup(html, "html.parser")
    for anchor in soup.find_all("a", href=True):
        text = _clean(anchor.get_text(" ", strip=True)).casefold()
        href = urljoin(LISTING_URL, anchor["href"])
        if "stáhnout kalendář" in text or "stahnout kalendar" in text:
            return href
        if href.casefold().endswith(".xlsx") and "kalend" in text:
            return href
    return None


def _xlsx_items(content: bytes) -> list[DiscoveryItem]:
    workbook = load_workbook(BytesIO(content), data_only=True)
    by_url: dict[str, DiscoveryItem] = {}

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows():
                row_title: str | None = None
                for cell in row:
                    if cell.value is not None and row_title is None:
                        row_title = _clean(str(cell.value))
                    target = None
                    if cell.hyperlink is not None:
                        target = cell.hyperlink.target
                    elif isinstance(cell.value, str) and cell.value.startswith(("http://", "https://")):
                        target = cell.value
                    if not target:
                        continue
                    url = urljoin(LISTING_URL, target)
                    if not _is_detail_url(url):
                        continue
                    by_url[url] = DiscoveryItem(
                        external_id=_external_id(url),
                        detail_url=url,
                        title_hint=row_title or PurePosixPath(urlsplit(url).path).name,
                        metadata={"discovery_method": "XLSX"},
                    )
    finally:
        workbook.close()

    return sorted(by_url.values(), key=lambda item: item.external_id)


class DotaceEuAdapter(SourceAdapter):
    descriptor = SourceDescriptor(
        code="DOTACEEU",
        name="DotaceEU.cz",
        authority=AuthorityLevel.OFFICIAL,
        country_code="CZ",
        base_url="https://www.dotaceeu.cz/",
        retrieval_modes=[RetrievalMode.HTML, RetrievalMode.XLSX],
        allowed_hosts=sorted(_ALLOWED_HOSTS),
        normal_refresh_minutes=240,
        max_concurrency=2,
        requests_per_second=0.5,
        disappearance_confirmation_runs=3,
        adapter_version="0.1.0",
    )

    async def healthcheck(self, ctx: AdapterContext) -> HealthReport:
        started = datetime.now(timezone.utc)
        try:
            response = await ctx.http.get(LISTING_URL)
            body = response.text.casefold()
            healthy = (
                response.status_code == 200
                and "výzvy" in body
                and "termín pro podání žádosti" in body
            )
            status = HealthStatus.HEALTHY if healthy else HealthStatus.DEGRADED
            detail = None if healthy else (
                f"Unexpected DotaceEU listing response: HTTP {response.status_code}"
            )
        except Exception as exc:
            status = HealthStatus.UNAVAILABLE
            detail = f"{type(exc).__name__}: {exc}"

        elapsed = datetime.now(timezone.utc) - started
        return HealthReport(
            status=status,
            checked_at=ctx.now,
            latency_ms=max(0, int(elapsed.total_seconds() * 1000)),
            detail=detail,
        )

    async def discover(
        self,
        ctx: AdapterContext,
        checkpoint: SourceCheckpoint | None,
    ) -> DiscoveryPage:
        del checkpoint
        response = await ctx.http.get(LISTING_URL)
        if response.status_code >= 400:
            raise RuntimeError(
                f"DotaceEU listing returned HTTP {response.status_code}"
            )

        listing_snapshot = self._snapshot(
            ctx,
            source_url=LISTING_URL,
            content=response.content,
            mime_type=response.headers.get(
                "content-type",
                "text/html",
            ).split(";", 1)[0],
            headers=response.headers,
        )

        html_items = _listing_html_items(response.text)
        items_by_url = {
            str(item.detail_url): item.model_copy(
                update={
                    "metadata": {
                        **item.metadata,
                        "listing_snapshot_id": listing_snapshot,
                    }
                }
            )
            for item in html_items
        }

        workbook_url = _xlsx_link(response.text)
        workbook_snapshot: str | None = None
        if workbook_url is not None:
            workbook_response = await ctx.http.get(workbook_url)
            if workbook_response.status_code < 400:
                workbook_snapshot = self._snapshot(
                    ctx,
                    source_url=workbook_url,
                    content=workbook_response.content,
                    mime_type=workbook_response.headers.get(
                        "content-type",
                        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    ).split(";", 1)[0],
                    headers=workbook_response.headers,
                )
                for item in _xlsx_items(workbook_response.content):
                    url = str(item.detail_url)
                    metadata = {
                        **item.metadata,
                        "listing_snapshot_id": listing_snapshot,
                        "calendar_snapshot_id": workbook_snapshot,
                        "calendar_url": workbook_url,
                    }
                    items_by_url[url] = item.model_copy(update={"metadata": metadata})

        items = sorted(items_by_url.values(), key=lambda item: item.external_id)
        return DiscoveryPage(
            items=items,
            next_checkpoint=None,
            is_complete=True,
            total_hint=len(items),
        )

    async def fetch_record(
        self,
        ctx: AdapterContext,
        item: DiscoveryItem,
        validators: FetchValidators | None = None,
    ) -> RecordFetchResult:
        headers: dict[str, str] = {}
        if validators:
            if validators.etag:
                headers["if-none-match"] = validators.etag
            if validators.last_modified:
                headers["if-modified-since"] = validators.last_modified

        response = await ctx.http.get(str(item.detail_url), headers=headers)
        if response.status_code == 304:
            return RecordFetchResult(
                state=FetchState.NOT_MODIFIED,
                http_status=304,
                etag=response.headers.get("etag"),
                last_modified=response.headers.get("last-modified"),
            )
        if response.status_code == 404:
            return RecordFetchResult(state=FetchState.GONE, http_status=404)
        if response.status_code >= 400:
            raise RuntimeError(
                f"DotaceEU detail returned HTTP {response.status_code}"
            )

        snapshot_id = self._snapshot(
            ctx,
            source_url=str(item.detail_url),
            content=response.content,
            mime_type=response.headers.get(
                "content-type",
                "text/html",
            ).split(";", 1)[0],
            headers=response.headers,
        )

        soup = BeautifulSoup(response.text, "html.parser")
        h1 = soup.find("h1")
        title = (
            _clean(h1.get_text(" ", strip=True))
            if h1
            else (item.title_hint or item.external_id)
        )
        pairs = _table_pairs(soup)

        call_code = _pair(pairs, "Číslo výzvy")
        call_type = _pair(pairs, "Druh výzvy")
        period = _pair(pairs, "Programové období")
        programme = _pair(pairs, "Operační program")
        priority_axis = _pair(pairs, "Prioritní osa")
        applicants = _pair(pairs, "Oprávnění žadatelé")
        access_text = _pair(pairs, "Zpřístupnění žádosti o podporu")
        opens_text = _pair(pairs, "Zahájení příjmu žádostí")
        closes_text = _pair(pairs, "Ukončení příjmu žádostí")
        status_text = _pair(pairs, "Stav výzvy")

        record = NativeRecord(
            source_code=self.descriptor.code,
            external_id=item.external_id,
            detail_url=item.detail_url,
            native_title=title,
            native_status=_status(status_text),
            raw_fields={
                "callCode": call_code,
                "callType": call_type,
                "programmingPeriod": period,
                "programme": programme,
                "priorityAxis": priority_axis,
                "eligibleApplicantsText": applicants,
                "applicationAccessAt": (
                    _parse_local_date(access_text).isoformat()
                    if _parse_local_date(access_text)
                    else None
                ),
                "submissionOpenAt": (
                    _parse_local_date(opens_text).isoformat()
                    if _parse_local_date(opens_text)
                    else None
                ),
                "submissionCloseAt": (
                    _parse_local_date(closes_text, end_of_day=True).isoformat()
                    if _parse_local_date(closes_text, end_of_day=True)
                    else None
                ),
                "officialStatusText": status_text,
                "moreInfoUrl": _more_info_url(soup),
                "history": _history_rows(soup),
                "discoveryMethod": item.metadata.get("discovery_method"),
            },
            snapshot_ids=[snapshot_id],
        )

        return RecordFetchResult(
            state=FetchState.MODIFIED,
            record=record,
            http_status=response.status_code,
            etag=response.headers.get("etag"),
            last_modified=response.headers.get("last-modified"),
        )

    async def fetch_artifact(
        self,
        ctx: AdapterContext,
        artifact: RemoteArtifactRef,
        validators: FetchValidators | None = None,
    ) -> ArtifactFetchResult:
        del ctx, artifact, validators
        raise NotImplementedError(
            "DotaceEU detail pages currently provide metadata/provenance; "
            "programme-specific official documents are collected by provider connectors."
        )

    def _snapshot(
        self,
        ctx: AdapterContext,
        *,
        source_url: str,
        content: bytes,
        mime_type: str,
        headers: Any,
    ) -> str:
        if ctx.snapshots is None:
            raise RuntimeError(
                "DotaceEU adapter requires ctx.snapshots for RAW-first ingestion"
            )
        snapshot = ctx.snapshots.put(
            source_code=self.descriptor.code,
            source_url=source_url,
            content=content,
            mime_type=mime_type,
            retrieved_at=ctx.now,
            validators={
                "etag": headers.get("etag") or "",
                "last_modified": headers.get("last-modified") or "",
            },
        )
        return snapshot.snapshot_id
