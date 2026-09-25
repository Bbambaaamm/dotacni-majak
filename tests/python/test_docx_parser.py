import io
import sys
import unittest
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pipelines" / "ingestion" / "src"))

from dotacni_majak_ingestion.document_model import BlockKind
from dotacni_majak_ingestion.document_security import DocumentSecurityError
from dotacni_majak_ingestion.docx_parser import (
    DocxDocumentError,
    DocxParserPolicy,
    parse_docx_document,
)


def _add_hyperlink(paragraph, text: str, url: str) -> None:
    relationship_id = paragraph.part.relate_to(
        url,
        RT.HYPERLINK,
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    run = OxmlElement("w:r")
    node = OxmlElement("w:t")
    node.text = text
    run.append(node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def _fixture_docx() -> bytes:
    document = Document()
    document.add_heading("Regiony 2026", level=1)
    document.add_paragraph("Podpora sportovni infrastruktury.")
    document.add_heading("Opravneni zadatele", level=2)
    document.add_paragraph("Obce a sportovni spolky.")

    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Parametr"
    table.cell(0, 1).text = "Hodnota"
    table.cell(1, 0).text = "Max. podpora"
    table.cell(1, 1).text = "90 %"

    paragraph = document.add_paragraph("Vice informaci: ")
    _add_hyperlink(paragraph, "podminky", "https://example.com/podminky")
    _add_hyperlink(paragraph, "unsafe", "javascript:alert(1)")

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


class DocxParserTest(unittest.TestCase):
    def test_extracts_headings_paragraphs_and_tables_with_stable_anchors(self):
        parsed = parse_docx_document(_fixture_docx())

        headings = [b for b in parsed.blocks if b.kind == BlockKind.HEADING]
        paragraphs = [b for b in parsed.blocks if b.kind == BlockKind.PARAGRAPH]
        tables = [b for b in parsed.blocks if b.kind == BlockKind.TABLE]

        self.assertEqual(headings[0].anchor.locator, "docx.heading.1")
        self.assertEqual(headings[0].anchor.section_path, ("Regiony 2026",))
        self.assertEqual(
            headings[1].anchor.section_path,
            ("Regiony 2026", "Opravneni zadatele"),
        )
        self.assertEqual(paragraphs[0].anchor.locator, "docx.p.1")
        self.assertEqual(tables[0].anchor.locator, "docx.table.1")
        self.assertEqual(
            tables[0].table_rows[1],
            ("Max. podpora", "90 %"),
        )
        self.assertIn("sportovni infrastruktury", parsed.canonical_text)

    def test_external_links_are_metadata_only_and_never_followed(self):
        parsed = parse_docx_document(_fixture_docx())
        self.assertEqual(
            parsed.metadata["external_links"],
            ("https://example.com/podminky",),
        )
        self.assertFalse(parsed.metadata["external_links_followed"])

    def test_rejects_corrupt_docx(self):
        with self.assertRaises(DocumentSecurityError):
            parse_docx_document(b"PK this is not a valid DOCX archive")

    def test_enforces_table_cell_limit(self):
        with self.assertRaises(DocxDocumentError):
            parse_docx_document(
                _fixture_docx(),
                policy=DocxParserPolicy(max_table_cells=2),
            )

    def test_enforces_text_limit(self):
        with self.assertRaises(DocxDocumentError):
            parse_docx_document(
                _fixture_docx(),
                policy=DocxParserPolicy(max_text_characters=10),
            )


if __name__ == "__main__":
    unittest.main()
