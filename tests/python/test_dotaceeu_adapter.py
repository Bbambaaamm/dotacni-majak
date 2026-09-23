import sys
import tempfile
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "connectors" / "dotaceeu" / "src"))
sys.path.insert(0, str(ROOT / "pipelines" / "ingestion" / "src"))

from dotacni_majak_dotaceeu import DotaceEuAdapter
from dotacni_majak_ingestion.snapshot import LocalRawSnapshotStore
from dotacni_majak_source_sdk import AdapterContext, GuardedHttpClient


LISTING = (
    ROOT / "connectors" / "dotaceeu" / "fixtures" / "listing.html"
).read_text(encoding="utf-8")
DETAIL = (
    ROOT / "connectors" / "dotaceeu" / "fixtures" / "call-109.html"
).read_text(encoding="utf-8")


def workbook_bytes() -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Výzvy"
    ws.append(["Název", "Detail"])
    ws.append([
        "MŽP_109. výzva",
        "Detail",
    ])
    cell = ws["B2"]
    cell.hyperlink = (
        "https://www.dotaceeu.cz/cs/jak-ziskat-dotaci/vyzvy/"
        "obdobi-2021-2027/05-operacni-program-zivotni-prostredi-2021-2027/"
        "mzp_109-vyzva,-sc-1-3,-opatreni-1-3-3,-prubezna"
    )
    ws.append(["IROP 120", "Detail"])
    ws["B3"].hyperlink = (
        "https://www.dotaceeu.cz/cs/jak-ziskat-dotaci/vyzvy/"
        "obdobi-2021-2027/06-integrovany-regionalni-operacni-program/"
        "120-vyzva-irop-kyberneticka-bezpecnost-ii"
    )
    data = BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


async def public_resolver(host: str, port: int) -> list[str]:
    return ["93.184.216.34"]


async def no_sleep(_: float) -> None:
    return None


class DotaceEuAdapterTest(unittest.IsolatedAsyncioTestCase):
    def handler(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path
        if path.rstrip("/") == "/cs/jak-ziskat-dotaci/vyzvy":
            return httpx.Response(
                200,
                text=LISTING,
                headers={"content-type": "text/html; charset=UTF-8"},
                request=request,
            )
        if path == "/media/calendar/vyzvy.xlsx":
            return httpx.Response(
                200,
                content=workbook_bytes(),
                headers={
                    "content-type": (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    )
                },
                request=request,
            )
        if "mzp_109-vyzva" in path:
            return httpx.Response(
                200,
                text=DETAIL,
                headers={"content-type": "text/html; charset=UTF-8"},
                request=request,
            )
        if "120-vyzva-irop" in path:
            return httpx.Response(
                200,
                text=DETAIL.replace(
                    "05_26_109",
                    "06_26_120",
                ).replace(
                    "MŽP_109. výzva, SC 1.3, opatření 1.3.3, průběžná",
                    "120. výzva IROP - Kybernetická bezpečnost II.",
                ),
                headers={"content-type": "text/html; charset=UTF-8"},
                request=request,
            )
        return httpx.Response(404, request=request)

    async def context(self, tmp):
        client = GuardedHttpClient(
            allowed_hosts={"dotaceeu.cz", "www.dotaceeu.cz"},
            resolver=public_resolver,
            sleeper=no_sleep,
            transport=httpx.MockTransport(self.handler),
        )
        ctx = AdapterContext(
            run_id="dotaceeu-test",
            http=client,
            logger=None,
            budget=None,
            snapshots=LocalRawSnapshotStore(tmp),
            now=datetime(2026, 9, 23, 12, 0, tzinfo=timezone.utc),
        )
        return client, ctx

    async def test_discovery_prefers_xlsx_hyperlinks_and_deduplicates_html(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                page = await DotaceEuAdapter().discover(ctx, None)
            finally:
                await client.aclose()

        self.assertTrue(page.is_complete)
        self.assertEqual(page.total_hint, 2)
        self.assertTrue(
            all(
                item.metadata["discovery_method"] == "XLSX"
                for item in page.items
            )
        )
        self.assertTrue(
            all("calendar_snapshot_id" in item.metadata for item in page.items)
        )

    async def test_detail_parses_official_fields_and_history(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                adapter = DotaceEuAdapter()
                page = await adapter.discover(ctx, None)
                item = next(
                    candidate
                    for candidate in page.items
                    if "mzp-109" in candidate.external_id
                )
                result = await adapter.fetch_record(ctx, item)
            finally:
                await client.aclose()

        record = result.record
        self.assertIsNotNone(record)
        self.assertEqual(record.native_status, "OPEN")
        self.assertEqual(record.raw_fields["callCode"], "05_26_109")
        self.assertEqual(record.raw_fields["callType"], "Průběžná")
        self.assertEqual(
            record.raw_fields["programme"],
            "Operační program Životní prostředí 2021—2027",
        )
        self.assertEqual(
            record.raw_fields["eligibleApplicantsText"],
            "Bez omezení, dle PrŽaP",
        )
        self.assertIsNotNone(record.raw_fields["submissionOpenAt"])
        self.assertIsNotNone(record.raw_fields["submissionCloseAt"])
        self.assertEqual(len(record.raw_fields["history"]), 2)
        self.assertEqual(
            record.raw_fields["moreInfoUrl"],
            "https://www.opzp.cz/dotace/109-vyzva/",
        )
        self.assertTrue(record.snapshot_ids)

    async def test_healthcheck_uses_public_listing(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                report = await DotaceEuAdapter().healthcheck(ctx)
            finally:
                await client.aclose()
        self.assertEqual(report.status.value, "HEALTHY")


if __name__ == "__main__":
    unittest.main()
