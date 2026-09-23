import sys
import tempfile
import unittest
from datetime import datetime, timezone
from io import BytesIO
from pathlib import Path

import httpx
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "connectors" / "irop" / "src"))
sys.path.insert(0, str(ROOT / "pipelines" / "ingestion" / "src"))

from dotacni_majak_ingestion.snapshot import LocalRawSnapshotStore
from dotacni_majak_irop import IropAdapter
from dotacni_majak_source_sdk import AdapterContext, GuardedHttpClient


LISTING = (
    ROOT / "connectors" / "irop" / "fixtures" / "listing.html"
).read_text(encoding="utf-8")
DETAIL = (
    ROOT / "connectors" / "irop" / "fixtures" / "call-120.html"
).read_text(encoding="utf-8")


def workbook_bytes() -> bytes:
    workbook = Workbook()
    ws = workbook.active
    ws.title = "Výzvy"
    ws.append(["Název", "Detail"])
    ws.append(["120. výzva IROP - Kybernetická bezpečnost II.", "Detail"])
    ws["B2"].hyperlink = (
        "https://irop.gov.cz/cs/vyzvy-2021-2027/vyzvy/120vyzvairop"
    )
    ws.append(["122. výzva IROP - Bezemisní vozidla", "Detail"])
    ws["B3"].hyperlink = (
        "https://irop.gov.cz/cs/vyzvy-2021-2027/vyzvy/122vyzvairop"
    )
    data = BytesIO()
    workbook.save(data)
    workbook.close()
    return data.getvalue()


async def public_resolver(host: str, port: int) -> list[str]:
    return ["93.184.216.34"]


async def no_sleep(_: float) -> None:
    return None


class IropAdapterTest(unittest.IsolatedAsyncioTestCase):
    def handler(self, request: httpx.Request) -> httpx.Response:
        path = request.url.path.rstrip("/")
        if path == "/cs/vyzvy-2021-2027":
            return httpx.Response(
                200,
                text=LISTING,
                headers={"content-type": "text/html; charset=UTF-8"},
                request=request,
            )
        if path == "/media/calendar/irop.xlsx":
            return httpx.Response(
                200,
                content=workbook_bytes(),
                headers={
                    "content-type": (
                        "application/vnd.openxmlformats-officedocument."
                        "spreadsheetml.sheet"
                    )
                },
                request=request,
            )
        if path.endswith("/120vyzvairop") or path.endswith("/122vyzvairop"):
            return httpx.Response(
                200,
                text=DETAIL,
                headers={"content-type": "text/html; charset=UTF-8"},
                request=request,
            )
        if path.startswith("/media/files/"):
            mime = "application/zip" if path.endswith(".zip") else "application/pdf"
            return httpx.Response(
                200,
                content=b"fixture",
                headers={"content-type": mime},
                request=request,
            )
        return httpx.Response(404, request=request)

    async def context(self, tmp):
        client = GuardedHttpClient(
            allowed_hosts={"irop.gov.cz", "www.irop.gov.cz"},
            resolver=public_resolver,
            sleeper=no_sleep,
            transport=httpx.MockTransport(self.handler),
        )
        return client, AdapterContext(
            run_id="irop-test",
            http=client,
            logger=None,
            budget=None,
            snapshots=LocalRawSnapshotStore(tmp),
            now=datetime(2026, 9, 23, 12, 0, tzinfo=timezone.utc),
        )

    async def test_discovery_uses_calendar_and_html_metadata(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                page = await IropAdapter().discover(ctx, None)
            finally:
                await client.aclose()

        self.assertEqual(page.total_hint, 3)
        ids = {item.external_id for item in page.items}
        self.assertIn("IROP-120", ids)
        self.assertIn("IROP-122", ids)
        i120 = next(i for i in page.items if i.external_id == "IROP-120")
        self.assertEqual(i120.native_status_hint, "OPEN")
        self.assertTrue(i120.metadata["calendar_snapshot_id"])

    async def test_detail_extracts_decision_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                adapter = IropAdapter()
                page = await adapter.discover(ctx, None)
                item = next(i for i in page.items if i.external_id == "IROP-120")
                record = (await adapter.fetch_record(ctx, item)).record
            finally:
                await client.aclose()

        self.assertIsNotNone(record)
        self.assertEqual(record.native_status, "OPEN")
        self.assertEqual(record.raw_fields["callNumber"], "120")
        self.assertEqual(record.raw_fields["callType"], "Průběžná")
        self.assertIn("kraje; obce", record.raw_fields["eligibleApplicantsText"])
        self.assertEqual(record.raw_fields["allocationCzk"], 1_798_341_944)
        self.assertEqual(len(record.raw_fields["history"]), 1)
        self.assertIn(
            "kybernetické bezpečnosti",
            record.raw_fields["generalInfoText"],
        )
        self.assertEqual(
            {artifact.role for artifact in record.artifacts},
            {"CALL_DOCUMENT", "GUIDELINES", "ANNEX"},
        )

    async def test_call_document_snapshot(self):
        with tempfile.TemporaryDirectory() as tmp:
            client, ctx = await self.context(tmp)
            try:
                adapter = IropAdapter()
                page = await adapter.discover(ctx, None)
                item = next(i for i in page.items if i.external_id == "IROP-120")
                record = (await adapter.fetch_record(ctx, item)).record
                artifact = next(
                    a for a in record.artifacts if a.role == "CALL_DOCUMENT"
                )
                result = await adapter.fetch_artifact(ctx, artifact)
            finally:
                await client.aclose()

        self.assertTrue(result.snapshot_id)
        self.assertGreater(result.size_bytes or 0, 0)


if __name__ == "__main__":
    unittest.main()
