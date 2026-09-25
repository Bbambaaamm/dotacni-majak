import io
import sys
import unittest
from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / "pipelines" / "ingestion" / "src"))

from dotacni_majak_ingestion.document_model import BlockKind
from dotacni_majak_ingestion.tabular_parser import (
    TabularDocumentError,
    TabularParserPolicy,
    parse_csv_document,
    parse_xlsx_document,
)


def _xlsx_fixture() -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Výzvy"
    sheet.append(["Název", "Podpora", "Datum", "Výpočet"])
    sheet.append(
        [
            "Sportovní infrastruktura",
            90.0,
            date(2026, 9, 25),
            "=B2/100",
        ]
    )
    sheet.merge_cells("A3:B3")
    sheet["A3"] = "Sloučená buňka"
    sheet["C3"] = datetime(2026, 10, 1, 12, 30)

    second = workbook.create_sheet("Skrytá data")
    second.sheet_state = "hidden"
    second.append(["Kód", "Hodnota"])
    second.append(["X", True])

    buffer = io.BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


class TabularParserTest(unittest.TestCase):
    def test_xlsx_preserves_row_provenance_and_value_types_as_text(self):
        parsed = parse_xlsx_document(_xlsx_fixture())

        self.assertEqual(parsed.format, "XLSX")
        self.assertTrue(all(b.kind == BlockKind.TABLE for b in parsed.blocks))
        first = parsed.blocks[0]
        second = parsed.blocks[1]
        self.assertEqual(first.anchor.locator, "xlsx.sheet.1.row.1")
        self.assertEqual(first.anchor.row, 1)
        self.assertEqual(first.anchor.section_path, ("Výzvy",))
        self.assertEqual(
            second.table_rows[0],
            (
                "Sportovní infrastruktura",
                "90",
                "2026-09-25",
                "=B2/100",
            ),
        )
        self.assertEqual(parsed.metadata["headers"]["Výzvy"][0], "Název")
        self.assertEqual(parsed.metadata["formula_cells"], 1)
        self.assertFalse(parsed.metadata["formulas_evaluated"])

    def test_xlsx_merged_cells_are_top_left_only_and_declared(self):
        parsed = parse_xlsx_document(_xlsx_fixture())
        merged = [
            b for b in parsed.blocks
            if b.anchor.locator == "xlsx.sheet.1.row.3"
        ][0]
        self.assertEqual(merged.table_rows[0][0], "Sloučená buňka")
        self.assertEqual(merged.table_rows[0][1], "")
        self.assertIn("A3:B3", parsed.metadata["merged_ranges"]["Výzvy"])
        self.assertEqual(parsed.metadata["merged_cell_policy"], "TOP_LEFT_ONLY")

    def test_hidden_sheet_is_retained_and_marked(self):
        parsed = parse_xlsx_document(_xlsx_fixture())
        self.assertEqual(parsed.metadata["sheet_states"]["Skrytá data"], "hidden")
        self.assertTrue(
            any(
                block.anchor.section_path == ("Skrytá data",)
                for block in parsed.blocks
            )
        )

    def test_csv_detects_semicolon_and_preserves_original_row_number(self):
        parsed = parse_csv_document(
            "\ufeffNázev;Podpora\nSportoviště;90\n\nDalší;80\n".encode("utf-8")
        )
        self.assertEqual(parsed.format, "CSV")
        self.assertEqual(parsed.metadata["delimiter"], ";")
        self.assertEqual(parsed.metadata["header"], ("Název", "Podpora"))
        self.assertEqual(parsed.blocks[0].anchor.locator, "csv.row.1")
        self.assertEqual(parsed.blocks[-1].anchor.locator, "csv.row.4")
        self.assertEqual(parsed.blocks[-1].anchor.row, 4)

    def test_csv_cp1250_fallback_is_explicit(self):
        parsed = parse_csv_document("Název;Město\nTest;Plzeň\n".encode("cp1250"))
        self.assertEqual(parsed.metadata["encoding"], "cp1250")
        self.assertIn("CSV_ENCODING_FALLBACK:cp1250", parsed.warnings)

    def test_xlsx_and_csv_limits_fail_closed(self):
        with self.assertRaises(TabularDocumentError):
            parse_xlsx_document(
                _xlsx_fixture(),
                policy=TabularParserPolicy(max_rows_per_sheet=1),
            )
        with self.assertRaises(TabularDocumentError):
            parse_csv_document(
                b"a;b;c\n1;2;3\n",
                policy=TabularParserPolicy(max_columns_per_sheet=2),
            )


if __name__ == "__main__":
    unittest.main()
