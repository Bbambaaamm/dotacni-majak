from __future__ import annotations

import csv
import io
import math
import re
from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal
from typing import Any

from .document_model import BlockKind, ParsedBlock, ParsedDocument, SourceAnchor
from .document_security import DocumentSecurityPolicy, inspect_document


class TabularDocumentError(ValueError):
    pass


class TabularDependencyMissingError(TabularDocumentError):
    pass


@dataclass(frozen=True, slots=True)
class TabularParserPolicy:
    max_input_bytes: int = 50 * 1024 * 1024
    max_sheets: int = 100
    max_rows_per_sheet: int = 250_000
    max_columns_per_sheet: int = 512
    max_total_cells: int = 5_000_000
    max_cell_characters: int = 100_000
    max_total_text_characters: int = 10_000_000

    def __post_init__(self) -> None:
        for name, value in (
            ("max_input_bytes", self.max_input_bytes),
            ("max_sheets", self.max_sheets),
            ("max_rows_per_sheet", self.max_rows_per_sheet),
            ("max_columns_per_sheet", self.max_columns_per_sheet),
            ("max_total_cells", self.max_total_cells),
            ("max_cell_characters", self.max_cell_characters),
            ("max_total_text_characters", self.max_total_text_characters),
        ):
            if value < 1:
                raise ValueError(f"{name} must be >= 1")


def _cell_text(
    value: Any,
    *,
    max_chars: int,
    excel_is_date: bool = False,
    excel_number_format: str | None = None,
) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        text = "TRUE" if value else "FALSE"
    elif isinstance(value, datetime):
        # openpyxl represents Excel date-only cells as midnight datetime.
        # Preserve date-only semantics when the cell is marked as a date and
        # its number format contains no explicit hour/second token.
        fmt = excel_number_format or ""
        has_explicit_time = bool(re.search(r"[hHsS]", fmt))
        if excel_is_date and value.time() == time(0, 0) and not has_explicit_time:
            text = value.date().isoformat()
        else:
            text = value.isoformat()
    elif isinstance(value, date):
        text = value.isoformat()
    elif isinstance(value, time):
        text = value.isoformat()
    elif isinstance(value, Decimal):
        text = format(value, "f")
    elif isinstance(value, float):
        if not math.isfinite(value):
            raise TabularDocumentError("non-finite numeric cell is not allowed")
        text = format(value, ".15g")
    else:
        text = str(value)

    text = " ".join(text.replace("\x00", " ").split())
    if len(text) > max_chars:
        raise TabularDocumentError("tabular cell text limit exceeded")
    return text


def _trim_trailing_empty(values: list[str]) -> tuple[str, ...]:
    while values and not values[-1]:
        values.pop()
    return tuple(values)


def _append_row_block(
    *,
    blocks: list[ParsedBlock],
    values: tuple[str, ...],
    locator: str,
    row_number: int,
    section_path: tuple[str, ...],
    state: dict[str, int],
    policy: TabularParserPolicy,
) -> None:
    if not values or not any(values):
        return

    state["cells"] += len(values)
    if state["cells"] > policy.max_total_cells:
        raise TabularDocumentError("tabular total cell limit exceeded")

    text = " | ".join(values)
    state["text_chars"] += len(text)
    if state["text_chars"] > policy.max_total_text_characters:
        raise TabularDocumentError("tabular total text limit exceeded")

    blocks.append(
        ParsedBlock(
            kind=BlockKind.TABLE,
            text=text,
            table_rows=(values,),
            anchor=SourceAnchor(
                locator=locator,
                row=row_number,
                section_path=section_path,
            ),
        )
    )


def parse_xlsx_document(
    content: bytes,
    *,
    policy: TabularParserPolicy | None = None,
) -> ParsedDocument:
    """Parse an XLSX without evaluating formulas or following external links."""

    policy = policy or TabularParserPolicy()
    inspect_document(
        content,
        declared_mime_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        policy=DocumentSecurityPolicy(max_file_bytes=policy.max_input_bytes),
    )

    try:
        from openpyxl import load_workbook
        from openpyxl.utils.exceptions import InvalidFileException
    except ImportError as exc:
        raise TabularDependencyMissingError(
            "XLSX parsing requires the ingestion 'documents' optional dependencies"
        ) from exc

    try:
        workbook = load_workbook(
            io.BytesIO(content),
            read_only=False,
            data_only=False,
            keep_links=False,
        )
    except (InvalidFileException, OSError, ValueError, KeyError) as exc:
        raise TabularDocumentError("invalid or corrupt XLSX") from exc
    except Exception as exc:
        raise TabularDocumentError("failed to open XLSX") from exc

    if len(workbook.worksheets) > policy.max_sheets:
        raise TabularDocumentError(
            f"XLSX sheet limit exceeded: {len(workbook.worksheets)} > {policy.max_sheets}"
        )

    blocks: list[ParsedBlock] = []
    state = {"cells": 0, "text_chars": 0}
    headers: dict[str, tuple[str, ...]] = {}
    merged_ranges: dict[str, tuple[str, ...]] = {}
    sheet_states: dict[str, str] = {}
    formula_cells = 0

    try:
        for sheet_index, worksheet in enumerate(workbook.worksheets, start=1):
            if worksheet.max_row > policy.max_rows_per_sheet:
                raise TabularDocumentError(
                    f"XLSX row limit exceeded in sheet {worksheet.title!r}: "
                    f"{worksheet.max_row} > {policy.max_rows_per_sheet}"
                )
            if worksheet.max_column > policy.max_columns_per_sheet:
                raise TabularDocumentError(
                    f"XLSX column limit exceeded in sheet {worksheet.title!r}: "
                    f"{worksheet.max_column} > {policy.max_columns_per_sheet}"
                )

            merged_ranges[worksheet.title] = tuple(
                sorted(str(value) for value in worksheet.merged_cells.ranges)
            )
            sheet_states[worksheet.title] = worksheet.sheet_state
            first_nonempty: tuple[str, ...] | None = None

            for row_number, row in enumerate(worksheet.iter_rows(), start=1):
                values: list[str] = []
                for cell in row:
                    if isinstance(cell.value, str) and cell.value.startswith("="):
                        formula_cells += 1
                    values.append(
                        _cell_text(
                            cell.value,
                            max_chars=policy.max_cell_characters,
                            excel_is_date=bool(getattr(cell, "is_date", False)),
                            excel_number_format=str(
                                getattr(cell, "number_format", "") or ""
                            ),
                        )
                    )

                frozen = _trim_trailing_empty(values)
                if frozen and any(frozen) and first_nonempty is None:
                    first_nonempty = frozen

                _append_row_block(
                    blocks=blocks,
                    values=frozen,
                    locator=f"xlsx.sheet.{sheet_index}.row.{row_number}",
                    row_number=row_number,
                    section_path=(worksheet.title,),
                    state=state,
                    policy=policy,
                )

            if first_nonempty is not None:
                headers[worksheet.title] = first_nonempty
    finally:
        workbook.close()

    return ParsedDocument(
        format="XLSX",
        blocks=tuple(blocks),
        canonical_text="\n".join(block.text for block in blocks),
        metadata={
            "sheets": tuple(sheet.title for sheet in workbook.worksheets),
            "headers": headers,
            "merged_ranges": merged_ranges,
            "merged_cell_policy": "TOP_LEFT_ONLY",
            "sheet_states": sheet_states,
            "formula_cells": formula_cells,
            "formulas_evaluated": False,
            "external_links_followed": False,
            "rows_emitted": len(blocks),
            "cells_emitted": state["cells"],
        },
    )


def _decode_csv(content: bytes) -> tuple[str, str | None]:
    try:
        return content.decode("utf-8-sig"), None
    except UnicodeDecodeError:
        try:
            return content.decode("cp1250"), "cp1250"
        except UnicodeDecodeError as exc:
            raise TabularDocumentError(
                "CSV is neither valid UTF-8 nor CP1250"
            ) from exc


def _csv_dialect(sample: str):
    try:
        return csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        class Fallback(csv.excel):
            delimiter = ";" if sample.count(";") > sample.count(",") else ","
        return Fallback


def parse_csv_document(
    content: bytes,
    *,
    policy: TabularParserPolicy | None = None,
) -> ParsedDocument:
    policy = policy or TabularParserPolicy()
    inspect_document(
        content,
        declared_mime_type="text/csv",
        policy=DocumentSecurityPolicy(max_file_bytes=policy.max_input_bytes),
    )

    text, fallback_encoding = _decode_csv(content)
    sample = text[:64_000]
    dialect = _csv_dialect(sample)
    reader = csv.reader(io.StringIO(text, newline=""), dialect=dialect)

    blocks: list[ParsedBlock] = []
    state = {"cells": 0, "text_chars": 0}
    header: tuple[str, ...] | None = None

    try:
        for row_number, row in enumerate(reader, start=1):
            if row_number > policy.max_rows_per_sheet:
                raise TabularDocumentError("CSV row limit exceeded")
            if len(row) > policy.max_columns_per_sheet:
                raise TabularDocumentError(
                    f"CSV column limit exceeded at row {row_number}"
                )

            values = _trim_trailing_empty(
                [
                    _cell_text(value, max_chars=policy.max_cell_characters)
                    for value in row
                ]
            )
            if values and any(values) and header is None:
                header = values

            _append_row_block(
                blocks=blocks,
                values=values,
                locator=f"csv.row.{row_number}",
                row_number=row_number,
                section_path=(),
                state=state,
                policy=policy,
            )
    except csv.Error as exc:
        raise TabularDocumentError("invalid CSV structure") from exc

    warnings = ()
    if fallback_encoding:
        warnings = (f"CSV_ENCODING_FALLBACK:{fallback_encoding}",)

    return ParsedDocument(
        format="CSV",
        blocks=tuple(blocks),
        canonical_text="\n".join(block.text for block in blocks),
        metadata={
            "delimiter": dialect.delimiter,
            "header": header or (),
            "encoding": fallback_encoding or "utf-8",
            "rows_emitted": len(blocks),
            "cells_emitted": state["cells"],
        },
        warnings=warnings,
    )
